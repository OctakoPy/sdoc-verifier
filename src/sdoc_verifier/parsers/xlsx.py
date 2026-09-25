"""Parser for Excel workbooks with label/value rows."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

from sdoc_verifier.models import Confidence, DocKind, ParsedDocument
from sdoc_verifier.parsers.txt import detect_kind

__all__ = ["parse_xlsx", "parse_xlsx_bytes"]


def _stringify(value: object) -> str:
    """Render a cell value the way the row contract expects.

    Floats are trimmed of a trailing ``.0`` (openpyxl reads ``15`` as
    ``15.0``); everything else is plain ``str``. ``None`` becomes ``""``.

    Example:
        >>> _stringify(15.0)
        '15'
        >>> _stringify("341715")
        '341715'
        >>> _stringify(None)
        ''
    """
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _kind_hint(rows: tuple[str, ...]) -> DocKind:
    """Document-kind hint from the banner/title rows of the grid.

    Scans the first few rows for a known banner phrase. In the working
    dataset the SI workbook's banner reads ``BL INSTRUCTION`` (it is an
    instruction *for* a bill of lading), while the BL workbook's banner
    reads ``BILL OF LADING`` — so ``BL INSTRUCTION`` maps to SI here even
    though a naive BL-substring rule would misfire on it.

    Example:
        >>> _kind_hint(("BL INSTRUCTION: 3154303911",))
        'SI'
        >>> _kind_hint(("BILL OF LADING: 3154303911",))
        'BL'
        >>> _kind_hint(("APRIL FINE PAPER TRADING",))
        'OTHER'
    """
    for row in rows[:4]:
        if "BL INSTRUCTION" in row.upper():
            return "SI"
        kind = detect_kind(row)
        if kind != "OTHER":
            return kind
    return "OTHER"


def parse_xlsx_bytes(path: str, data: bytes) -> tuple[ParsedDocument, Confidence]:
    """Parse raw ``.xlsx`` bytes into a document + confidence pair.

    Args:
        path: Attachment path exactly as listed in the email JSON.
        data: Raw file bytes (the caller read them from disk).

    Returns:
        ``(document, confidence)`` with ``fmt="xlsx"``. Corrupt or empty
        workbooks come back ``kind=UNREADABLE`` with confidence 0.1 and an
        ``error`` string instead of raising.
    """
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(BytesIO(data), data_only=True, read_only=True)
    except Exception as exc:  # noqa: BLE001 - any workbook error is "unreadable"
        document = ParsedDocument(
            path=path,
            fmt="xlsx",
            kind="UNREADABLE",
            text="",
            rows=(),
            scanned=False,
            error=f"cannot read workbook: {exc}",
        )
        return document, Confidence(0.1, "corrupt or unsupported xlsx")

    try:
        lines: list[str] = []
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows(values_only=True):
                cells = [_stringify(cell) for cell in row]
                while cells and not cells[-1]:
                    cells.pop()
                if not cells:
                    continue
                if len(cells) == 1:
                    lines.append(cells[0])
                else:
                    lines.append(f"{cells[0]}: {' | '.join(cells[1:])}")
    finally:
        workbook.close()

    text = "\n".join(lines)
    if not text.strip():
        document = ParsedDocument(
            path=path,
            fmt="xlsx",
            kind="UNREADABLE",
            text="",
            rows=(),
            scanned=False,
            error="workbook has no cell content",
        )
        return document, Confidence(0.1, "empty sheet")

    rows = tuple(lines)
    kind = _kind_hint(rows)
    document = ParsedDocument(
        path=path,
        fmt="xlsx",
        kind=kind,
        text=text,
        rows=rows,
        scanned=False,
        error=None,
    )
    return document, Confidence(
        0.9, f"parsed grid: {len(rows)} row(s), kind hint {kind}"
    )


def parse_xlsx(path: str | Path) -> tuple[ParsedDocument, Confidence]:
    """Parse one ``.xlsx`` attachment from disk.

    Missing files return an ``UNREADABLE`` document (confidence 0.1) rather
    than raising, mirroring the txt parser's error contract.

    Args:
        path: Filesystem path of the attachment.

    Returns:
        ``(document, confidence)`` as in :func:`parse_xlsx_bytes`.

    Example:
        >>> import tempfile, os
        >>> from openpyxl import Workbook
        >>> wb = Workbook(); ws = wb.active
        >>> ws["A1"] = "SHIPPER"; ws["B1"] = "ACME LTD"
        >>> with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as fh:
        ...     wb.save(fh.name); tmp = fh.name
        >>> doc, conf = parse_xlsx(tmp)
        >>> (doc.fmt, doc.rows)
        ('xlsx', ('SHIPPER: ACME LTD',))
        >>> os.unlink(tmp)
    """
    source = Path(path)
    try:
        data = source.read_bytes()
    except OSError as exc:
        document = ParsedDocument(
            path=str(path),
            fmt="xlsx",
            kind="UNREADABLE",
            text="",
            rows=(),
            scanned=False,
            error=f"cannot read file: {exc}",
        )
        return document, Confidence(0.1, "unreadable file")
    return parse_xlsx_bytes(str(path), data)
