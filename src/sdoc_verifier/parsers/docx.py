"""Parser for Word documents with paragraph and table content."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

from sdoc_verifier.models import Confidence, ParsedDocument
from sdoc_verifier.parsers.txt import detect_kind

__all__ = ["parse_docx", "parse_docx_bytes"]


def _fold_value(value: str) -> str:
    """Fold a multi-line cell value onto one line (``; `` separators).

    Example:
        >>> _fold_value("ACME LTD\\n77 ROBINSON ROAD\\nSINGAPORE")
        'ACME LTD; 77 ROBINSON ROAD; SINGAPORE'
    """
    parts = [part.strip() for part in value.splitlines() if part.strip()]
    return "; ".join(parts)


def _kind_hint(lines: tuple[str, ...]) -> str:
    """Document-kind hint from the leading paragraph lines.

    Uses the shared txt-header detector (``BILL OF LADING (DRAFT)`` → BL,
    ``SHIPPING INSTRUCTION`` → SI). ``B/L NO.(提单号)`` meta lines are
    deliberately not kind signals — they can appear on either document.

    Example:
        >>> _kind_hint(("BILL OF LADING (DRAFT)", "B/L NO.(提单号): X"))
        'BL'
    """
    for line in lines[:6]:
        kind = detect_kind(line)
        if kind != "OTHER":
            return kind
    return "OTHER"


def parse_docx_bytes(path: str, data: bytes) -> tuple[ParsedDocument, Confidence]:
    """Parse raw ``.docx`` bytes into a document + confidence pair.

    Args:
        path: Attachment path exactly as listed in the email JSON.
        data: Raw file bytes.

    Returns:
        ``(document, confidence)`` with ``fmt="docx"``. Corrupt or empty
        files come back ``kind=UNREADABLE`` with confidence 0.1 and an
        ``error`` string instead of raising.
    """
    from docx import Document as read_document

    try:
        document = read_document(BytesIO(data))
        paragraphs = tuple(
            para.text.strip() for para in document.paragraphs if para.text.strip()
        )
        rows: list[str] = []
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                while cells and not cells[-1]:
                    cells.pop()
                if not cells:
                    continue
                if len(cells) == 1:
                    rows.append(cells[0])
                else:
                    rows.append(f"{cells[0]}: {_fold_value(cells[1])}")
    except Exception as exc:  # noqa: BLE001 - any docx error is "unreadable"
        failed = ParsedDocument(
            path=path,
            fmt="docx",
            kind="UNREADABLE",
            text="",
            rows=(),
            scanned=False,
            error=f"cannot read docx: {exc}",
        )
        return failed, Confidence(0.1, "corrupt or unsupported docx")

    lines = paragraphs + tuple(rows)
    if not lines:
        failed = ParsedDocument(
            path=path,
            fmt="docx",
            kind="UNREADABLE",
            text="",
            rows=(),
            scanned=False,
            error="docx has no paragraph or table content",
        )
        return failed, Confidence(0.1, "empty document")

    kind = _kind_hint(paragraphs)
    text = "\n".join(lines)
    parsed = ParsedDocument(
        path=path,
        fmt="docx",
        kind=kind,  # type: ignore[arg-type]  # hint narrowed below at extract
        text=text,
        rows=lines,
        scanned=False,
        error=None,
    )
    return parsed, Confidence(
        0.9, f"parsed docx: {len(paragraphs)} para, {len(rows)} table rows"
    )


def parse_docx(path: str | Path) -> tuple[ParsedDocument, Confidence]:
    """Parse one ``.docx`` attachment from disk.

    Missing files return an ``UNREADABLE`` document (confidence 0.1) rather
    than raising, mirroring the other parsers' error contract.

    Args:
        path: Filesystem path of the attachment.

    Returns:
        ``(document, confidence)`` as in :func:`parse_docx_bytes`.
    """
    source = Path(path)
    try:
        data = source.read_bytes()
    except OSError as exc:
        failed = ParsedDocument(
            path=str(path),
            fmt="docx",
            kind="UNREADABLE",
            text="",
            rows=(),
            scanned=False,
            error=f"cannot read file: {exc}",
        )
        return failed, Confidence(0.1, "unreadable file")
    return parse_docx_bytes(str(path), data)
